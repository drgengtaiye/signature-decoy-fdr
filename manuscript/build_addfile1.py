"""Additional file 1 for the split Genome Biology manuscript.

usage: build_addfile1.py <Additional_file_1_Notes_and_Figures.docx> <Additional_file_1_figures_pdf dir>
                         <FigS_ISPY2.pdf> <FigS_ISPY2.png> <outdir>

Works on the GM version element by element (indices below are its body children):
drops the attribution, triage, locked-prediction, TME-library and external-validation notes and
figures S5-S12, adds the I-SPY2 figure (old main Fig. 8g-i) as S5, renumbers old S13 to S6, and
remaps main-figure numbers (old 4-6 -> 5-7).
"""
import copy
import re
import shutil
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Emu
from docx.text.paragraph import Paragraph

SRC, PDFDIR, SPDF, SPNG, OUT = sys.argv[1:6]
OUT = Path(OUT)
d = docx.Document(SRC)
E = list(d.element.body.iterchildren())
P = lambda i: Paragraph(E[i], d)
text = lambda el: "".join(t.text or "" for t in el.iter(qn("w:t")))

# guard: the indices below belong to the GM version
for i, start in [(14, "The locked predictions"), (17, "The holdouts"), (18, "The attribution at protein"),
                 (20, "Attribution against"), (21, "Attribution and triage"), (22, "Gene weighting"),
                 (23, "Two questions"), (24, "A library the model"), (25, "Against tumor mutational"),
                 (26, "External validation"), (29, "All changes below"), (45, "Fig. S4."), (48, "Fig. S5."),
                 (69, "Fig. S12."), (72, "Fig. S13.")]:
    assert text(E[i]).startswith(start), (i, text(E[i])[:40])


def rich(p, s):
    """Rewrite paragraph p from marked-up text: *...* is italic; a bare P before =, < or ≤ is italic."""
    runs = p.runs
    r_norm = next(r for r in runs if not r.italic)._r.find(qn("w:rPr"))
    r_ital = next(r for r in runs if r.italic)._r.find(qn("w:rPr"))
    for r in runs:
        r._r.getparent().remove(r._r)
    segs = []
    for k, part in enumerate(re.split(r"\*([^*]+)\*", s)):
        if k % 2:
            segs.append((part, True))
            continue
        for j, q in enumerate(re.split(r"\bP(?= [=<≤])", part)):
            if j:
                segs.append(("P", True))
            segs.append((q, False))
    for t, it in segs:
        if not t:
            continue
        r = p.add_run(t)
        rp = r_ital if it else r_norm
        if rp is not None:
            r._r.insert(0, copy.deepcopy(rp))


def sub(i, old, new):
    """Replace text inside one run of element i."""
    for r in Paragraph(E[i], d).runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return
    raise AssertionError((i, old))


# ---- front
sub(0, "Figures S1–S13", "Figures S1–S6")
p1 = P(1)
for r in p1.runs[1:]:
    r._r.getparent().remove(r._r)
p1.runs[0].text = "Matched decoy modules give immune gene-expression signatures a false discovery rate"
sub(33, "Figures S1–S13", "Figures S1–S6")

# ---- Note S1
sub(16, "Fig. S13c", "Fig. S6c")
rich(P(22), "*A null without simulation.* The closed-form null predicts the drawn null’s median to 1.6% and its "
     "95th percentile to 1.9% (median absolute error, leave-one-cohort-out), its standard deviation to 8.5%, and "
     "changes 38 of 630 signature × cohort calls (6.0%), above the 5% the plan allowed; the drawn null therefore "
     "stays, and the simulation is not an optional part of the protocol.")

# the two pre-specified questions (Methods) used seven cohorts collected for the split-off validation;
# they stay reported here, with the cohorts named
rich(P(23), "*Two questions that did not resolve.* Both were fixed in advance (Methods) and both use, besides the 12 "
     "checkpoint cohorts, seven checkpoint-treated cohorts collected for a separate analysis and not used elsewhere "
     "in this paper: GEO GSE176307 (urothelial), GSE161537 (non-small-cell lung), GSE159067 (head and neck), "
     "GSE215868 (melanoma), GSE93157 (several tumor types), GSE207422 (non-small-cell lung, with chemotherapy) and "
     "GSE235919 (colorectal, with chemotherapy). The spread of the inflamed score across a cohort’s patients ranges "
     "from 0.40 to 0.72 and correlates with the cohort’s mean signature effect at rho = +0.18 (19 cohorts, "
     "one-sided P = 0.23; +0.05 with the renal cohorts removed, +0.23 given cohort size). Whatever makes a cohort "
     "able to show the effect, it is not captured by how much immune expression varies there. The first principal "
     "component of the 52 scores carries 47% of their variance and reaches an AUC of 0.53 to 0.85 across the seven "
     "further cohorts; against APM (Thompson), the signature with the largest mean effect in the 12 checkpoint "
     "cohorts, it is +0.004 higher on average but with a one-sided 95% lower bound of −0.075, so non-inferiority at "
     "the pre-specified margin of 0.02 is not established. That comparison rests on the four cohorts where APM can "
     "also be scored, and the axis ranks 14th of 44 signatures in the median cohort.")

# ---- Note S2: keep changes (i), (ii) and (vi); (iii)-(v) concern the split-off analyses
rich(P(29), "All changes below were made before the plan naming the third randomized trial was deposited, and every "
     "result reported uses the final version. (i) Three signatures (Tcell-inflamed, Ayers; CTL, Jiang; "
     "WNT-β-catenin, Takeuchi) had been lost when the curation’s metadata file was parsed; they were restored and "
     "every benchmark recomputed. (ii) Nulls were first drawn from 48-gene sets scored as mean z-scores on both "
     "sides; they were replaced by nulls matched on size and engine (Fig. 1b, c). Against the fixed null the "
     "inflamed family cleared 52% of response tests, against 32% under the matched null; on the prognostic side "
     "matching raised the mean gain in C-index from +0.008 to +0.018. (iii) The interaction null was first run on "
     "a single score subspace and is now reported as the weakest over K = 3 to 20, and the renal contrast was "
     "first tested by patient-level logistic regression (P = 0.008) before the cohort-level permutation was "
     "adopted. The target-decoy analysis was designed and run after the analyses above were complete; the directed "
     "statistic, the 1,000-draw null banks and the natural gene pool it uses are post hoc and are marked as such "
     "where they appear.")

# ---- legends kept: remap main-figure numbers, drop the attribution clause
sub(39, "Figs 3a and 6g", "Figs 3a and 7g")
sub(42, ": the two-compartment division that the single-cell attribution explains (Fig. S5).", ".")
sub(45, "Fig. 4e.", "Fig. 5e.")
sub(72, "Fig. S13.", "Fig. S6.")

# ---- I-SPY2 as Fig. S5: copy the S4 block (blank, image, legend) and swap image and legend
blank, img_t, leg_t = (copy.deepcopy(E[i]) for i in (43, 44, 45))
cx = int(next(E[44].iter(qn("wp:extent"))).get("cx"))
tmp = d.add_paragraph()
tmp.add_run().add_picture(SPNG, width=Emu(cx))
drawing_new = next(tmp._p.iter(qn("w:drawing")))
drawing_old = next(img_t.iter(qn("w:drawing")))
drawing_old.getparent().replace(drawing_old, drawing_new)
tmp._p.getparent().remove(tmp._p)
leg = Paragraph(leg_t, d)
bold_rpr = leg.runs[0]._r.find(qn("w:rPr"))
for r in leg.runs:
    r._r.getparent().remove(r._r)
r = leg.add_run("Fig. S5. The third randomized trial, I-SPY2.")
r._r.insert(0, copy.deepcopy(bold_rpr))
# the body of the legend takes the normal and italic run formats of the S4 legend
leg_runs_src = Paragraph(E[45], d).runs
norm_rpr = next(x for x in leg_runs_src if not x.bold and not x.italic)._r.find(qn("w:rPr"))
ital_rpr = next(x for x in leg_runs_src if x.italic)._r.find(qn("w:rPr"))
BODY = (" (a) Interaction of each of the 17 immune (blue) and 3 control (vermilion) gene sets of the trial library "
        "with arm on pathological complete response (pCR ~ score + arm + score × arm, logistic; score standardized "
        "within the contrast), for paclitaxel plus pembrolizumab (filled; 69 against 179 patients) and durvalumab "
        "plus olaparib (open; 71 against 29) against paclitaxel alone; bars, 95% confidence intervals. All 17 immune "
        "interactions favor pembrolizumab and none reaches P < 0.05 (pre-specified count, P = 1.0 against the "
        "no-interaction simulation; Methods). (b) Pembrolizumab contrast: each of the 52 signatures’ AUC for "
        "pathological complete response in the paclitaxel-only arm against the pembrolizumab arm, colored by family; "
        "dotted, identity. The inflamed/IFN signatures discriminate in both arms (median AUC 0.73 and 0.66), which "
        "leaves little room for an interaction. (c) All four randomized comparisons: immune gene sets (of 17) whose "
        "interaction favors the checkpoint arm (bar) and that reach P < 0.05 in that direction (diamond): "
        "IMmotion150 17 and 8, PRINCE 14 and 5, I-SPY2 pembrolizumab 17 and 0, durvalumab plus olaparib 3 and 0.")
for j, q in enumerate(re.split(r"\bP(?= [=<≤])", BODY)):
    for t, rp in ([("P", ital_rpr)] if j else []) + [(q, norm_rpr)]:
        rr = leg.add_run(t)
        if rp is not None:
            rr._r.insert(0, copy.deepcopy(rp))

# ---- drop: Note S1 paragraphs on the split-off analyses,
# and figure blocks S5-S12 (elements 46-69)
DROP = [14, 17, 18, 20, 21, 24, 25, 26] + list(range(46, 70))
anchor = E[45]
for el in (blank, img_t, leg_t):
    anchor.addnext(el); anchor = el
for i in DROP:
    E[i].getparent().remove(E[i])

# the images of the deleted figures are still related to the document part; drop them
body_xml = d.element.xml
for rid, rel in list(d.part.rels.items()):
    if rel.reltype.endswith("/image") and f'r:embed="{rid}"' not in body_xml:
        d.part.drop_rel(rid)                   # unreferenced only (drop_rel would also drop a count of 1)

OUT.mkdir(parents=True, exist_ok=True)
out = OUT / "Additional_file_1_Notes_and_Figures.docx"
d.save(out)

# ---- vector PDFs, one per figure
pdf_out = OUT / "Additional_file_1_figures_pdf"
pdf_out.mkdir(exist_ok=True)
for new, src in [(1, "S1"), (2, "S2"), (3, "S3"), (4, "S4"), (6, "S13")]:
    shutil.copy2(Path(PDFDIR) / f"Supplementary_Figure_{src}.pdf", pdf_out / f"Supplementary_Figure_S{new}.pdf")
shutil.copy2(SPDF, pdf_out / "Supplementary_Figure_S5.pdf")

# ---- checks
d2 = docx.Document(out)
alltext = "\n".join(p.text for p in d2.paragraphs)
heads = re.findall(r"^(Fig\. S\d+)\.", alltext, re.M)
assert heads == [f"Fig. S{i}" for i in range(1, 7)], heads
bad = re.findall(r".{0,60}(attribut|triage|locked|Fig\. [78]|Figs? S(?:[7-9]|1\d)|external validation|validation cohort|atlas).{0,40}",
                 alltext, re.I)
print("figures:", heads, "| images:", len(list(d2.element.body.iter(qn("w:drawing")))))
print("residual terms:", bad if bad else "none")
print("main-figure refs:", sorted(set(re.findall(r"Figs?\.? \d+[a-z]?(?:,? ?(?:and )?\d*[a-z])*", alltext))))
print("words:", len(alltext.split()))
