"""MANUSCRIPT_GB_v1 -> v2 text fixes, before the reference rebuild (build_refs_GB.py runs after this).

usage: build_manuscript_GB_v2.py <MANUSCRIPT_GB_v1.docx> <out.docx>

- Fig. 2 legend for the redrawn figure (fig2_gb.py: no atlas / triage rings, no alluvial)
- Additional file 1 references: S<I-SPY2> -> S5, S13 -> S6; its entry in the list; Add. file 3 entry
- validation / development wording left over from the split, and the seven further cohorts of the
  two pre-specified questions named by accession
- code paths that pointed into the repository layout now point at Additional files 2 and 3
"""
import copy
import re
import sys

import docx
from docx.oxml.ns import qn

SRC, OUT = sys.argv[1:3]
d = docx.Document(SRC)
PS = d.paragraphs


def para(start):
    hits = [p for p in PS if p.text.startswith(start)]
    assert len(hits) == 1, (start, len(hits))
    return hits[0]


def sub(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return
    raise AssertionError((p.text[:40], old))


def rewrite(p, body, lead=None):
    """Replace the text of p; a bold lead run (legend title) is kept as the first run if given."""
    runs = p.runs
    keep = 1 if lead else 0
    if lead:
        assert runs[0].text.strip() == lead.strip(), runs[0].text
    tmpl = runs[keep]._r.find(qn("w:rPr"))
    for r in runs[keep:]:
        r._r.getparent().remove(r._r)
    r = p.add_run(body)
    if tmpl is not None:
        r._r.insert(0, copy.deepcopy(tmpl))


n_sub = 0
for p in PS:
    for r in p.runs:
        for old, new in [("Fig. S‹I-SPY2›", "Fig. S5"), ("Fig. S13", "Fig. S6")]:
            if old in r.text:
                r.text = r.text.replace(old, new); n_sub += 1
assert n_sub == 4, n_sub
assert "‹" not in "".join(p.text for p in PS)

# ---- Fig. 2 legend
p = para("Figure 2. The 52 signatures")
lead = p.runs[0].text
old = p.text[len(lead):]
b_part = old[old.index(" (b) Shared genes"):old.index(" (c) Effective number")]
d_part = old[old.index(" (d) Correlation between"):old.index(" (e) The 50 attributed")]
rewrite(p, " (a) Every signature as concentric tracks, grouped by family and ordered by response effect. From the "
        "center: family; scoring engine; gene count (log scale); response, prognostic and shift effects in null s.d. "
        "(outward, positive; gray, negative); declared direction (black, higher in responders, 34 signatures; gray, "
        "lower, 18), the direction in which the directed statistic scores each signature." + b_part +
        " (c) Effective number of independent tests among the 52 in each checkpoint cohort (Li and Ji; patients in "
        "parentheses; color, tumor type); dashed, median 17 (range 15–23)." + d_part.rstrip() + " ", lead=lead)

# ---- additional files list
p = para("Additional file 1. Supplementary Notes")
rewrite(p, " Supplementary Notes S1–S2 and Figures S1–S6. File name: Additional_file_1_Notes_and_Figures.docx; "
        "format: Word document (.docx). Note S1: alternative tests and sensitivity analyses. Note S2: changes made "
        "during the analysis. Supplementary figures with their legends; the same figures are supplied as vector "
        "PDFs, one per file, in Additional_file_1_figures_pdf/. Fig. S1: The shift does not depend on how signatures "
        "were scored. Fig. S2: The gastric discovery cohort in detail. Fig. S3: The gastric cohorts at closer "
        "inspection. Fig. S4: The randomized trials in detail. Fig. S5: The third randomized trial, I-SPY2. "
        "Fig. S6: The information ceiling. ", lead="Additional file 1.")
p = para("Additional file 3. Analysis code.")
sub(p, "Scripts that build every derived table, figure and manuscript version, and the validation code.",
    "Scripts that build every derived table and figure and the manuscript, including the target-decoy pipeline, "
    "with the analysis plans as deposited and their SHA-256 hashes.")

# ---- leftovers of the split
p = para("In each of the 12 development cohorts the inflamed signatures")
sub(p, "In each of the 12 development cohorts", "In each of the 12 checkpoint cohorts")
sub(p, "icb/positive2.R", "positive2.R, Additional file 3")

p = para("A third plan fixed three questions before any was computed")
rewrite(p, "A third plan fixed three questions before any was computed; the third is the gene-level test above. "
        "The first asked whether a cohort’s own spread of immune expression explains how much signatures gain "
        "there: in each of the 12 checkpoint cohorts and seven further checkpoint-treated cohorts collected for a "
        "separate analysis (Additional file 1: Note S1), the inflamed/interferon family’s genes, standardized across "
        "patients, were averaged per patient and the standard deviation of that score correlated with the cohort’s "
        "mean effect over the 52 signatures (one-sided Spearman). The second asked, in the seven further cohorts, "
        "whether the library is one axis: the 52 standardized scores’ first principal component, its sign fixed by "
        "the inflamed family and no outcome used, against the signature with the largest mean effect in the 12 "
        "checkpoint cohorts, fixed in advance; the endpoint was the paired difference in AUC with a cluster "
        "bootstrap over cohorts and a non-inferiority margin of 0.02. Neither question resolved (Additional file 1: "
        "Note S1). ")

p = para("Counts of significant signatures were referred")
sub(p, "Changes made to the analysis before the external validation was specified are listed in",
    "Changes made during the analysis, and which analyses are post hoc, are listed in")

p = para("IMmotion150 randomized to atezolizumab alone")
sub(p, "pancan/v4_arm_split.csv", "Additional file 2, sheet v4_arm_split")
sub(p, "icb/arm_split.py", "arm_split.py, Additional file 3")

p = para("Association with response in treated patients")
sub(p, "and a third, with pathological complete response, was added at validation.",
    "and a third, with pathological complete response, was found later by a pre-specified search.")
p = para("PRINCE randomized patients with pancreatic cancer")
sub(p, "the third trial is reported with the validation below.", "the third trial is reported below.")

p = para("All data analyzed in this study are publicly available")
txt = p.text
anchor = "and GSE173839 [88]."
assert anchor in txt
for r in p.runs:
    if anchor in r.text:
        r.text = r.text.replace(anchor, anchor + " Seven further checkpoint-treated cohorts, used only for two "
                                "pre-specified questions (Additional file 1: Note S1), are GEO accessions GSE176307, "
                                "GSE161537, GSE159067, GSE215868, GSE93157, GSE207422 and GSE235919.")
        break
else:
    # the anchor spans runs: append the sentence after the run that ends it
    end = txt.index(anchor) + len(anchor); pos = 0
    for r in p.runs:
        pos += len(r.text)
        if pos >= end:
            assert pos == end, "anchor ends inside a run"
            new = copy.deepcopy(r._r); r._r.addnext(new)
            for t in new.iter(qn("w:t")): t.text = ""
            next(new.iter(qn("w:t"))).text = (" Seven further checkpoint-treated cohorts, used only for two "
                "pre-specified questions (Additional file 1: Note S1), are GEO accessions GSE176307, GSE161537, "
                "GSE159067, GSE215868, GSE93157, GSE207422 and GSE235919.")
            next(new.iter(qn("w:t"))).set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            break

# ==== Genome Biology format (Research article guidelines): abstract <= 250 words, Methods after
# Conclusions, use of an LLM documented in Methods, tables numbered in order of first citation

# ---- abstract: 249 words; the body runs are not bold (v1 inherited bold from the label run)
ABSTRACT = {
    "Background: ": "A gene-expression signature is usually declared to work when it beats random gene sets of the same "
                   "size, but nothing says how many such claims in a library are false. Proteomics solved the same "
                   "problem with target-decoy searches.",
    "Results: ": "We grew 2,912 decoy modules from genes in no published signature, each matched to one of 52 immune "
                "signatures in size, scoring engine, expression level and internal co-expression across 12 "
                "checkpoint-treated cohorts (827 patients). The published analysis, a folded AUC against a size- and "
                "engine-matched null, called 21 signatures at an estimated false discovery rate (FDR) of 17% (95% CI "
                "11–27%), although its per-test error was nominal: co-expression widens the spread of a module's AUC, "
                "folding reads its favourable tail, and pooling cohorts turns the excess into calls. Scoring each "
                "signature in its declared direction called 27 at an FDR of 7–14%, depending on how closely decoys "
                "were matched, with the gain in weakly co-expressed signatures; folding had counted as a success a "
                "signature that runs significantly the wrong way. Strongly co-expressed decoys passed at 2.6 times "
                "the nominal rate. In two randomized trials, signatures beat their null in 15 of 34 checkpoint-arm "
                "tests and 0 of 34 comparator-arm tests (pooled hazard ratio 0.66 per standard deviation, one-sided "
                "P = 0.0026).",
    "Conclusions: ": "Matched decoys attach a false discovery rate to a claim about a signature library and expose "
                    "errors that per-test calibration misses. Signatures should be scored in their declared direction "
                    "and judged against decoys that share their co-expression.",
}
nwords = 0
for lab, body in ABSTRACT.items():
    p = para(lab)
    lab_run = p.runs[0]
    assert lab_run.text == lab, lab_run.text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    parts = re.split(r"\bP(?= = )", body)
    for j, q in enumerate(parts):
        for t, it in ([("P", True)] if j else []) + [(q, False)]:
            r = p.add_run(t)
            rpr = lab_run._r.find(qn("w:rPr"))
            if rpr is not None:
                r._r.insert(0, copy.deepcopy(rpr))
            r.bold = False
            r.italic = it or None
    nwords += len(p.text.split())
assert nwords <= 250, nwords
print("abstract words:", nwords)


def heading(text, style_id):
    els = [el for el in d.element.body.iterchildren()
           if el.tag == qn("w:p") and "".join(t.text or "" for t in el.iter(qn("w:t"))).strip() == text
           and el.find(qn("w:pPr")) is not None and el.find(qn("w:pPr")).find(qn("w:pStyle")) is not None
           and el.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val")) == style_id]
    assert len(els) == 1, (text, len(els))
    return els[0]


# ---- use of a large language model, as the last Methods subsection (wording for the authors to confirm)
last_sub = heading("A third randomized trial", "5")
last_body = last_sub.getnext()
assert "".join(t.text or "" for t in last_body.iter(qn("w:t"))).startswith("A third randomized trial was identified")
h = copy.deepcopy(last_sub)
for t in h.iter(qn("w:t")): t.text = ""
next(h.iter(qn("w:t"))).text = "Use of a large language model"
b = copy.deepcopy(last_body)
for r in list(b.iter(qn("w:r")))[1:]:
    r.getparent().remove(r)
for t in b.iter(qn("w:t")): t.text = ""
next(b.iter(qn("w:t"))).text = ("Claude (Anthropic), used through the Claude Code command-line tool, assisted in writing and "
                                "running analysis code, assembling figures, checking references against PubMed, and "
                                "drafting and editing the text. The authors directed the work, checked the code, "
                                "results and text, and take full responsibility for the content.")
last_body.addnext(h); h.addnext(b)

# ---- move Methods (heading through the last subsection) to after Conclusions
body = d.element.body
kids = list(body.iterchildren())
i0 = kids.index(heading("Methods", "4")); i1 = kids.index(heading("Results", "4"))
i_ab = kids.index(heading("Abbreviations", "4"))
assert i0 < i1 < kids.index(heading("Conclusions", "4")) < i_ab
anchor = kids[i_ab]
for el in kids[i0:i1]:
    anchor.addprevious(el)

# ---- tables in order of first citation: the claims table is now cited first
SWAP = {"Table 1": "Table 2", "Table 2": "Table 1"}
n_tab = 0
for p in d.paragraphs:
    for r in p.runs:
        if re.search(r"Table [12]\b", r.text):
            r.text = re.sub(r"Table ([12])\b", lambda m: SWAP[m.group(0)], r.text); n_tab += 1
assert n_tab == 4, n_tab

# ---- declarations: GB's list has no clinical-trial heading; its competing-interests wording
ct = heading("Clinical trial information", "5")
na = ct.getnext()
assert "".join(t.text or "" for t in na.iter(qn("w:t"))).strip() == "Not applicable."
ct.getparent().remove(ct); na.getparent().remove(na)
p = para("The authors declare no competing interests.")
sub(p, "The authors declare no competing interests.", "The authors declare that they have no competing interests.")

d.save(OUT)

# ---- checks
d2 = docx.Document(OUT)
alltxt = "\n".join(p.text for p in d2.paragraphs)
body = alltxt[:alltxt.index("Figure titles and legends")]
for pat in [r"attribut", r"triage", r"atlas attribution", r"external validation", r"validation cohort",
            r"development (cohort|effect)", r"Fig\. S‹", r"Figs?\. S(?:[7-9]|1\d)", r"icb/|pancan/"]:
    hits = [m.group(0) for m in re.finditer(r".{0,50}" + pat + r".{0,30}", alltxt, re.I)]
    print(f"{pat}: {hits if hits else 'none'}")
print("Fig. S refs:", sorted(set(re.findall(r"Figs?\. S\d+[a-z]?", alltxt))))
