"""Integrated Genome Biology manuscript, version 1: target-decoy FDR as the spine, attribution and
triage split off into their own paper.

    python build_manuscript_GB.py <MANUSCRIPT_GM1015.docx> <outdir>
        -> MANUSCRIPT_GB_v1.docx, CHANGELOG_GB_v1.md

Works on the submitted manuscript itself, so its equations, styles and untouched paragraphs carry
over. Paragraph numbers below are those of MANUSCRIPT_GM1015.docx. New references are written as
{Author Year} because the bibliography is to be rebuilt from PMIDs (REFERENCES_audit.csv); the old
[n] numbers are left as they are until then.
"""
import copy, os, re, sys
import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import build_targetdecoy as TD
import build_front_v2 as FR
SRC, OUTD = sys.argv[1], sys.argv[2]
I = "/groups/adv2105_gp/yichen/Yi/com/gastric/icb/"
MAT = pd.read_csv(I + "calib_v3_fdr_matched.csv").set_index(["null", "stat", "subset"])
LOG = []

d = Document(SRC)
P = list(d.paragraphs); E = [p._p for p in P]
assert P[143].text.startswith("In 12 checkpoint-treated cohorts") and P[183].text.startswith("This study has further")
ORIG_TEXT = {i: p.text for i, p in enumerate(P)}
has_math = lambda el: "oMath" in el.xml

def first_rpr(el):
    r = el.find(qn("w:r"))
    rp = r.find(qn("w:rPr")) if r is not None else None
    return copy.deepcopy(rp) if rp is not None else None

def fill(el, parts):
    """Replace a paragraph's content by runs; parts = [(text, bold)], formatting of its first run kept."""
    rpr = first_rpr(el)
    for ch in list(el):
        if ch.tag != qn("w:pPr"): el.remove(ch)
    par = Paragraph(el, d)
    for text, bold in parts:
        r = par.add_run(text)
        if rpr is not None: r._r.insert(0, copy.deepcopy(rpr))
        if bold: r.bold = True
    return el

def set_text(i, text, bold_prefix=None, why=""):
    if has_math(E[i]): LOG.append(f"- note: paragraph {i} contains an equation and was replaced whole ({why})")
    fill(E[i], ([(bold_prefix, True)] if bold_prefix else []) + [(text, False)])

def new_after(anchor, like, text, bold_prefix=None):
    el = copy.deepcopy(E[like]); fill(el, ([(bold_prefix, True)] if bold_prefix else []) + [(text, False)])
    anchor.addnext(el); return el

def delete(*idx):
    for i in idx:
        if E[i].getparent() is not None: E[i].getparent().remove(E[i])

def cite_set(text):
    out = set()
    for grp in re.findall(r"\[([\d,\s–\-]+)\]", text):
        for part in grp.split(","):
            part = part.strip()
            if re.match(r"^\d+[–\-]\d+$", part):
                a, b = map(int, re.split(r"[–\-]", part)); out |= set(range(a, b + 1))
            elif part.isdigit(): out.add(int(part))
    return out
REF0 = next(i for i, p in enumerate(P) if p.text.strip() == "References")
cited_before = set().union(*(cite_set(P[i].text) for i in range(REF0))) | \
    set().union(*(cite_set(c.text) for t in d.tables for r in t.rows for c in r.cells))

# ---------------------------------------------------------------- front matter
set_text(0, FR.TITLE)
for i, (k, v) in zip((10, 11, 12), FR.AB.items()): set_text(i, v, bold_prefix=k)
delete(13); set_text(15, FR.KW)

# ---------------------------------------------------------------- Background
bg = Document(os.path.join(HERE, "..", "deliverables", "REVISED_background_paragraph.docx"))
pos = next(p.text for p in bg.paragraphs if p.text.startswith("How a gene set should be compared"))
NEWREF = {9: "Goeman 2007", 10: "Subramanian 2005", 11: "Tamayo 2016", 12: "Wu 2012", 13: "Wu 2010",
          14: "Barry 2005", 15: "Efron 2007"}
pos = re.sub(r"\[([\d,\s]+)\]", lambda m: "{" + "; ".join(NEWREF[int(x)] for x in m.group(1).split(",")) + "}", pos)
new_after(E[18], 18, pos)
delete(20)
set_text(21,
 "Here we score 52 published signatures by each publication’s own method and compare each with random gene sets "
 "matched to it on size and scoring engine, for response in 12 checkpoint-treated cohorts (827 patients) and for "
 "survival in 20 TCGA cancer types (7,091 patients). We then ask how many of the resulting claims are false. "
 "Benchmarks of published signatures have counted how many beat random gene sets [6] {Beck 2013; Waldron 2014; "
 "Tang 2017}, and one estimated a false discovery rate for the internal coherence of 12,096 published signatures "
 "from random signatures drawn from signature genes {Cantini 2018}; none has put an error rate on the count of "
 "signatures that predict an outcome, against comparison sets matched to each signature. "
 "Borrowing the target-decoy design of peptide identification {Elias 2007}, we give each signature decoy modules "
 "matched to it in size, scoring engine, expression level and internal co-expression, put signatures and decoys "
 "through the identical test, and read a false discovery rate for the library’s claims off the decoys. Because the "
 "response and survival settings differ in endpoint, population and assay as well as in treatment, we finally test "
 "the treatment interaction within randomized trials.")

# ---------------------------------------------------------------- Methods
set_text(28,
 "The response and prognostic benchmarks, the target-decoy calibration, the dissociation and the library overview "
 "use all 52 signatures. The randomized analyses use one gene-set library of 31 signatures whose genes are present "
 "in both trial matrices, 17 of them immune and 3 confounder controls; ESTIMATE immune [11], produced by a separate "
 "package, is not part of it. No subset was chosen by its result.")
delete(62, 63, 64)                                             # triage ridge model
t66 = ORIG_TEXT[66].replace(" and 12 for the 50 attributed ones", "")
assert t66 != ORIG_TEXT[66]; set_text(66, t66)
t90 = ORIG_TEXT[90]
cut90 = t90.index("All twelve plans are in the code archive")
set_text(90, t90[:cut90] + "The plans are in the code archive with their hashes and deposit times.")
set_text(85, "The survival effect")
# paragraph 86: keep the survival sentence and its equations, cut the attribution null that follows
p86 = E[86]; cut = False
for ch in list(p86):
    if ch.tag == qn("w:pPr"): continue
    if cut: p86.remove(ch); continue
    t = "".join(x.text or "" for x in ch.iter(qn("w:t")))
    if ch.tag == qn("w:r") and "The attribution carries" in t:
        keep = t.split("The attribution carries")[0].rstrip()
        ts = list(ch.iter(qn("w:t"))); ts[0].text = keep
        for x in ts[1:]: x.getparent().remove(x)
        cut = True
assert cut, "paragraph 86 not cut"
delete(87, 88)
tdh = new_after(E[55], 47, "Target-decoy calibration")
_m = TD.METHODS.replace("Target-decoy calibration. ", "", 1)
_a = "The expected number of false discoveries is"
assert _m.count(_a) == 1
_m = _m.replace(_a, "As in a separate target-decoy search with several decoys per target {Käll 2008; Keich 2019}, the "
                    "expected number of false discoveries is")
_b = "(FDR) is that sum divided by the number of targets called."
assert _m.count(_b) == 1
_m = _m.replace(_b, _b + " The FDR is estimated, not controlled: unlike knockoffs {Barber 2015}, decoys carry no "
                    "guarantee of exchangeability with the false targets.")
new_after(tdh, 48, _m)
delete(*range(111, 130))                                       # attribution, triage, refinement, learned weights
set_text(130, "A third randomized trial")
set_text(131,
 "A third randomized trial was identified by a search of GEO on 2026-09-18 through NCBI E-utilities for human "
 "expression series mentioning a checkpoint antibody or checkpoint blockade, under an analysis plan deposited before "
 "any cohort it named was scored. I-SPY2 (GSE194040) [86, 87] contributed the HER2-negative paclitaxel controls and "
 "the paclitaxel plus pembrolizumab arm, and GSE173839 [88] the durvalumab plus olaparib arm and its controls, on the "
 "deposited batch-corrected gene-level arrays. Each score entered pCR ~ score + arm + score:arm (logistic), and the "
 "count of checkpoint-favoring interactions at P < 0.05 among the 17 immune gene sets was referred to a "
 "no-interaction simulation drawing pathological complete response from the fitted pCR ~ PC1..PC9 + arm model "
 "(2,000 replicates).")
delete(132, 133, 134, 135)                                     # second atlas, TRIPOD+AI reporting

# ---------------------------------------------------------------- Results
h = new_after(E[143], 142, TD.RES[0]); prev = h
for k, txt in enumerate(TD.RES[1:]):
    prev = new_after(prev, 143 if k == 0 else 146, txt)
t172 = (ORIG_TEXT[172]
        .replace("The search also yielded a third randomized trial, I-SPY2, in which HER2-negative early breast cancer "
                 "was randomized", "A third randomized trial, I-SPY2, found by a pre-specified search (Methods), "
                 "randomized HER2-negative early breast cancer")
        .replace("(Fig. 8g, h)", "(Additional file 1: Fig. S‹I-SPY2›)").replace("(Fig. 8i)", "(Additional file 1: Fig. S‹I-SPY2›)"))
assert "‹I-SPY2›" in t172 and "found by a pre-specified search" in t172
set_text(172, t172); el = E[172]; el.getparent().remove(el); E[154].addnext(el)
LOG.append("- I-SPY2 paragraph (old paragraph 172) moved to the end of the randomized-trial section; old Fig. 8g–i to become a figure of Additional file 1, placeholder `Fig. S‹I-SPY2›`")
delete(*range(164, 172), 173, 174)                             # attribution, triage, locked predictions, TME library, lung atlas

# Table 2: drop the two attribution rows, add the target-decoy claim
t2 = d.tables[1]; rows = t2.rows
assert rows[6].cells[0].text.startswith("The cells that produce") and rows[7].cells[0].text.startswith("The attribution is more")
for r in (rows[7], rows[6]): r._tr.getparent().remove(r._tr)
m0, m1 = MAT.loc[("N0", "fold", "decoys within 0.05 r")], MAT.loc[("N0", "dir", "decoys within 0.05 r")]
r0, r1 = TD.r0, TD.r1
new_tr = copy.deepcopy(t2.rows[1]._tr); t2.rows[1]._tr.addnext(new_tr)
TDROW = ["Claims about the library carry a measurable false discovery rate",
         f"{len(TD.INFO):,} decoy modules matched to the 52 signatures in size, engine, expression level and internal "
         "co-expression, through the identical test",
         f"Published folded analysis: {TD.called('N0','fold')} called, FDR {r0.fdr:.0f}% ({r0.fdr_lo:.0f}–{r0.fdr_hi:.0f}%); "
         f"directed: {TD.called('N0','dir')} called, FDR {r1.fdr:.0f}% ({r1.fdr_lo:.0f}–{r1.fdr_hi:.0f}%)",
         f"With only the closest-matched decoys the directed FDR is {m1.fdr:.0f}% (published {m0.fdr:.0f}%); strongly "
         "co-expressed decoys carry response information of their own; one application"]
from docx.table import _Row
row = _Row(new_tr, t2)
for c, txt in zip(row.cells, TDROW):
    for p in c.paragraphs[1:]: p._p.getparent().remove(p._p)
    fill(c.paragraphs[0]._p, [(txt, False)])

# ---------------------------------------------------------------- Discussion and Conclusions
new_after(E[176], 177,
 "A claim that published signatures work is a claim about a library, and it can carry a false discovery rate. "
 "Decoy modules matched to each of the 52 signatures put that number on the analysis this field uses, and that our own "
 f"first analysis used: of the {TD.called('N0','fold')} signatures a folded AUC calls against a size- and "
 f"engine-matched null, about one in {["", "", "two", "three", "four", "five", "six", "seven", "eight", "nine", "ten"][round(100 / r0.fdr)]} is expected to be false, although the null is "
 "calibrated test by test. The excess comes from folding, which reads the wider spread of a co-expressed module at "
 "its favourable tail, and from pooling cohorts, which turns a small consistent excess into a call. Scoring in the "
 "declared direction finds more signatures; how much it lowers the error depends on co-expression, a large reduction "
 "for weakly co-expressed signatures and none for strongly co-expressed ones, whose decoys pass often because "
 "coherent gene programs carry response information whatever their content. The decoys also exposed a flaw in one of "
 "our own matched nulls, drawn from a pool that over-represented signature genes, which no per-test check had caught. "
 "Proteomics adopted target-decoy estimation because it attaches an error rate to a list of claims without a model of "
 "the null {Elias 2007}; gene-expression signatures are in the same position, and decoys can be built for any "
 "library, scoring rule and endpoint. To our knowledge this is the first application of target-decoy estimation, with "
 "decoys matched to each signature on scoring rule, weights, expression level and internal co-expression, to a claim "
 "about the predictive performance of published signatures. A false discovery rate has been estimated for the "
 "internal coherence of a large signature collection, from random signatures drawn from the signatures’ own genes "
 "{Cantini 2018}; for an outcome, we find, a pool that over-represents signature genes biases the null.")
delete(181)                                                    # attribution paragraph
set_text(182,
 "Four recommendations follow. Immune signatures intended for immunotherapy should be evaluated against treatment "
 "response and, where a randomized comparison exists, against the treatment interaction. A signature should be scored "
 "in its declared direction: folding the AUC counts wrong-way signatures as successes and inflates false claims. A "
 "claim about a library, that so many of its signatures work, should be reported with a false discovery rate, which "
 "matched decoy modules provide and per-test calibration of the null does not, and co-expressed signatures should be "
 "judged against decoys that share their co-expression. And any new signature should clear a random-gene-set null "
 "matched on size and scoring engine and a proliferation/hypoxia control, against which prognostic claims reported "
 "without these controls should be re-examined before they are taken at face value.")
set_text(183,
 "This study has further limitations. The decoys are matched on four properties, not on everything: they fall short "
 "of the internal co-expression of the most tightly co-expressed signatures, and restricting the estimate to the "
 f"closest-matched decoys moves the directed FDR from {r1.fdr:.0f}% to {m1.fdr:.0f}% while leaving the published "
 f"analysis at {m0.fdr:.0f}%. A decoy is null only in carrying no gene of any of the 52 signatures; strongly "
 "co-expressed decoys carry response information of their own, so for coherent signatures the FDR measures what a "
 "signature’s content adds beyond its coherence, which is the question a claim of specificity raises but not the "
 "only one. The estimate comes from one application, response to checkpoint blockade in 12 cohorts, and its "
 "intervals are wide. The prognostic analysis is within cohort in TCGA, with cross-cohort validation only in gastric "
 "cancer. Six signatures requiring bespoke algorithms could not be scored, the checkpoint cohorts are weighted toward "
 "melanoma, and response is dichotomized RECIST. The within-arm comparison is descriptive, and the arms differ in "
 "size: matching them attenuates it, from 9 and 6 of 17 to a median of 4 of 17 in each trial.")
set_text(185,
 "A claim that a gene-expression signature works can carry a false discovery rate. Decoy modules matched to each of "
 "52 published immune signatures show that the usual analysis of checkpoint response, a folded AUC against a "
 f"size-matched null, called {TD.called('N0','fold')} at an FDR of about {r0.fdr:.0f}%, an error that per-test "
 "calibration does not reveal, and that scoring in the declared direction calls more signatures. What the signatures "
 "measure is specific to the treatment given: inside randomized trials they beat matched random gene sets in the "
 "checkpoint arms and not in the comparator arms, so survival in checkpoint-naive cohorts is a weak test of a "
 "signature meant for immunotherapy. Signatures should be scored in their declared direction, judged against "
 "treatment response and, where possible, the treatment interaction, and claims about a library should be reported "
 "with a decoy-based false discovery rate.")


# ---------------------------------------------------------------- Declarations
set_text(196,
 "All data analyzed in this study are publicly available and no new sequencing data were generated. Bulk expression "
 "and clinical data for 20 TCGA cancer types were obtained from UCSC Xena [66], with survival endpoints from the TCGA "
 "Clinical Data Resource [67]. Harmonized immune-checkpoint cohorts were obtained from the ORCESTRA clinical_icb "
 "collection [63] (deposited on Zenodo) as MultiAssayExperiment objects. IMmotion150 and PRINCE were obtained from "
 "cBioPortal [94], and the I-SPY2 data are GEO accessions GSE194040 [86, 87] and GSE173839 [88]. Published signature "
 "definitions were obtained from bhklab/SignatureSets [4]. Gastric bulk validation cohorts are TCGA-STAD [68], "
 "GSE84437 [69] and GSE62254 [70], and the gastric checkpoint cohort is ENA accession PRJEB25780 [54]. Derived tables "
 "underlying every figure are provided as Additional file 2, and the decoy modules with their test results as "
 "Additional file 4. The analysis plans are deposited with the code, each with its hash and deposit time; Methods "
 "states for each whether it was written before or after the results it builds on.")
set_text(197,
 "All analysis code is provided as Additional file 3 and in the immunoTriage-paper repository [95], including the "
 "scripts that grow and score the decoy modules and estimate the false discovery rate (calib_decoys_v3.R, "
 "calib_score_v3.R, null_natpool_v4.R, calib_v3_fdr.py, calib_v3_robust.py). Signature scoring by each publication’s "
 "own method and the size- and engine-matched null are provided in the R package immunoTriage (version 0.3.1) [96]. "
 "Both repositories will be archived with a DOI on acceptance.")

# ---------------------------------------------------------------- legends and additional files
delete(317, 319)                                               # Figures 7 and 8 (attribution, locked predictions)
leg = new_after(E[309], 309, TD.LEG[1], bold_prefix=TD.LEG[0] + " ")
NEWNUM = {E[196], E[197], leg}                                 # written in the new numbering already
delete(322, 323)                                               # locked triage predictions, TRIPOD+AI checklist
af4 = new_after(E[325], 325,
 "Additional file 4. Target-decoy tables. File name: Additional_file_4_target_decoy.xlsx; format: Excel workbook "
 f"(.xlsx). The {len(TD.INFO):,} decoy modules (target, pool, structure, size, engine, declared direction and matching "
 "diagnostics), the test result of every signature and decoy under each null and statistic, the false discovery "
 "rates, the paired comparisons and the matching sensitivity analyses.")
LOG.append("- the Additional file 1 description (figures S1–S13) is still the old one: the attribution figures among S5–S12 go, S‹I-SPY2› comes in; redo with Additional file 1")

# abbreviations: add FDR, drop the ones the split paper no longer uses
_bt = [p.text for p in d.paragraphs]
_ri, _li = _bt.index("References"), _bt.index("Figure titles and legends")
_main = "\n".join(t for t in _bt[:_ri] + _bt[_li:] if not t.startswith("ACRG:"))
_main += "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
lines = [l for l in ORIG_TEXT[187].split("\n") if l.strip()] + ["FDR: false discovery rate"]
keep = [l for l in lines if re.search(r"(?<![A-Za-z])" + re.escape(l.split(":")[0]) + r"(?![A-Za-z])", _main)]
gone = [l.split(":")[0] for l in lines if l not in keep]
if gone: LOG.append("- abbreviations dropped as no longer used in the text: " + ", ".join(gone))
fill(E[187], [("\n".join(sorted(keep, key=lambda s: s.lower())), False)])

# ---------------------------------------------------------------- renumbering
FIG = {4: 5, 5: 6, 6: 7}
ADD = {4: 2, 5: 3}
rx_fig = re.compile(r"\b(Figs?\.\s?|Figures?\s)(\d+)(?!\d)")    # "Fig. 4d" has no word boundary after the 4
rx_add = re.compile(r"\b(Additional file\s)(\d+)\b")
rx_afn = re.compile(r"\b(Additional_file_)(\d+)(_)")
def renum(text, skip_new=False):
    text = rx_fig.sub(lambda m: m.group(1) + str(FIG.get(int(m.group(2)), int(m.group(2)))), text)
    text = rx_add.sub(lambda m: m.group(1) + str(ADD.get(int(m.group(2)), int(m.group(2)))), text)
    return rx_afn.sub(lambda m: m.group(1) + str(ADD.get(int(m.group(2)), int(m.group(2)))) + m.group(3), text)
def all_pars():
    for p in d.paragraphs: yield p
    for t in d.tables:
        for r in t.rows:
            for c in r.cells: yield from c.paragraphs
seen = set()
for p in all_pars():
    if p._p in seen: continue
    seen.add(p._p)
    if p._p in NEWNUM or p._p is af4: continue
    old = p.text; want = renum(old)
    if want == old: continue
    for r in p.runs: r.text = renum(r.text)
    if p.text != want:
        if has_math(p._p): LOG.append(f"- figure number not changed automatically (spans runs and contains an equation): {old[:80]}…")
        else: fill(p._p, [(want, False)]); LOG.append(f"- figure number spans runs, paragraph rewritten: {old[:60]}…")
# the new figure is Figure 4
for p in all_pars():
    if "Fig. X" in p.text or "Figure X" in p.text:
        for r in p.runs: r.text = r.text.replace("Fig. X", "Fig. 4").replace("Figure X", "Figure 4")

# ---------------------------------------------------------------- checks
body = [p.text for p in d.paragraphs]
ri = body.index("References")
text_main = "\n".join(body[:ri]) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
for pat in (r"Fig\. ?[89]\b", r"Figure [89]\b", r"\b[Tt]riage\b", r"\b[Aa]ttribut", r"\batlas\b", r"\blocked\b",
            r"Additional file [5-9]\b", r"TRIPOD"):
    for m in re.finditer(pat, text_main):
        s = text_main[max(0, m.start() - 70): m.end() + 50].replace("\n", " ")
        LOG.append(f"- residual `{pat}`: …{s}…")
cited_after = cite_set(text_main)
dropped = sorted(cited_before - cited_after)
words = lambda a, b: sum(len(t.split()) for t in body[body.index(a) + 1: body.index(b)])
out = os.path.join(OUTD, "MANUSCRIPT_GB_v1.docx"); d.save(out)
with open(os.path.join(OUTD, "CHANGELOG_GB_v1.md"), "w") as fh:
    fh.write("# MANUSCRIPT_GB_v1 change log (written by scripts/build_manuscript_GB.py)\n\n")
    fh.write(f"Words: Background {words('Background','Methods')}, Methods {words('Methods','Results')}, "
             f"Results {words('Results','Discussion')}, Discussion {words('Discussion','Conclusions')}\n\n")
    fh.write("## For manual attention\n\n" + "\n".join(LOG) + "\n\n")
    fh.write("## Old reference numbers no longer cited after the split\n\n" + ", ".join(map(str, dropped)) +
             "\n\n(drop when the reference list is rebuilt; new references are {author year} placeholders: Elias 2007, Goeman 2007, Subramanian 2005, Tamayo 2016, "
             "Wu 2012, Wu 2010, Barry 2005, Efron 2007, Cantini 2018, Beck 2013, Waldron 2014, Tang 2017, Käll 2008, Keich 2019, "
             "Barber 2015; PMIDs/DOIs in PRIORART_20261019.md and REVISED_background_paragraph.docx)\n")
print("saved", out); print("\n".join(LOG)); print("dropped refs:", dropped)
