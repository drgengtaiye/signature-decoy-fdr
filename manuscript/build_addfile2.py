"""Additional file 2 (derived tables) for the Genome Biology split.

usage: build_addfile2.py <Additional_file_4_derived_tables.xlsx from the GM package> <outdir>

Copies the GM workbook with openpyxl and removes the sheets that backed the attribution,
triage, locked-prediction / external-validation and TME-library analyses, which moved to the
separate attribution paper. Kept sheets are untouched (values and formats), except the two I-SPY2
interaction sheets, which also carried the separate paper's locked-signature library and keep only
the trial library; the key sheet loses the rows of the removed sheets and gets the new row counts.
"""
import sys
from pathlib import Path

import openpyxl

SRC, OUTD = Path(sys.argv[1]), Path(sys.argv[2])
OUT = OUTD / "Additional_file_2_derived_tables.xlsx"

# sheet -> why it leaves (everything else stays)
DROP = {
    "singlecell_attribution": "single-cell attribution",
    "spatial_attribution": "spatial attribution",
    "gene_attribution": "gene-level attribution",
    "v4_module_attribution": "module attribution",
    "attribution_v2_null": "attribution null",
    "proteomics_attribution": "protein-level attribution",
    "atlas_independence_results": "second-atlas test of the attribution",
    "v3_sigtriage_final": "triage calls",
    "v3_triage_loso": "triage model, leave-one-signature-out",
    "v3_shift_loso": "triage model, shift prediction",
    "v3_holdouts": "triage holdouts",
    "v3_atlas_transfer": "atlas transfer for triage",
    "v3_atlas_transfer_bycohort": "atlas transfer for triage",
    "v3_barrier_triage": "triage of the gastric scores",
    "v4sens_holdouts": "triage sensitivity analysis",
    "v4sens_sigtriage_final": "triage sensitivity analysis",
    "refine_test": "attribution-guided gene removal",
    "gw_test_results": "attribution-guided gene weighting",
    "validation_results": "external validation of the locked triage calls",
    "validation_by_signature": "external validation of the locked triage calls",
    "TME_LOCKED_predictions": "locked calls for the TME library",
    "tme_response_matched": "TME-library validation",
    "tme_overlap_with_52": "TME-library validation",
    "arrangement_robust": "dropped spatial-arrangement strand",
}

wb = openpyxl.load_workbook(SRC)
missing = set(DROP) - set(wb.sheetnames)
assert not missing, f"sheets not in workbook: {missing}"
for s in DROP:
    wb.remove(wb[s])

# I-SPY2 interaction sheets: drop the rows of the separate paper's locked-signature library
TRIAL_ONLY = ["ispy2_interaction_by_signature", "ispy2_interaction_summary"]
for s in TRIAL_ONLY:
    ws = wb[s]
    lib = [c.value for c in ws[1]].index("library") + 1
    for r in range(ws.max_row, 1, -1):
        if ws.cell(r, lib).value != "trial library":
            ws.delete_rows(r)

key = wb["key"]
hdr = [c.value for c in key[1]]
col = hdr.index("sheet") + 1
for r in range(2, key.max_row + 1):
    if key.cell(r, col).value in TRIAL_ONLY:
        key.cell(r, hdr.index("rows") + 1).value = wb[key.cell(r, col).value].max_row - 1
for r in range(key.max_row, 1, -1):
    if key.cell(r, col).value in DROP:
        key.delete_rows(r)
listed = [key.cell(r, col).value for r in range(2, key.max_row + 1)]
data = [s for s in wb.sheetnames if s != "key"]
assert listed == data, (listed, data)

OUTD.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"{OUT}: kept {len(data)} data sheets, dropped {len(DROP)}")
