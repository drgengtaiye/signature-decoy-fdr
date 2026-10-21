#!/usr/bin/env python3
"""Renumber the citations of the split GB manuscript and rebuild its reference list.

usage: python3 -I build_refs_GB.py <in.docx> <out.docx> [<REFERENCES_corrected.docx> <newrefs dir>]

Citations are numbered in reading order: body paragraphs and table cells in document order,
then the figure legends and additional-file entries after the reference list. Numeric
citations ([6], [43, 44], [54-63]) and {Author Year; Author Year} placeholders are both read;
adjacent citations ("[6] {Beck 2013; Tang 2017}") are merged into one bracket. Old entries
are taken from the corrected author lists (REFERENCES_corrected.docx, numbering as in the
GM manuscript); placeholders are resolved against the PubMed / Crossref records saved in
source/newrefs/ (esummary.json, efetch.xml, crossref_*.json), keyed by first-author surname
and year. Edits are made on w:t elements, so run formatting and OMML are left alone.
"""
import copy
import csv
import json
import re
import sys
import unicodedata
import xml.etree.ElementTree as ET
from pathlib import Path

import docx
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SRC, OUT = sys.argv[1], sys.argv[2]
CORR = sys.argv[3] if len(sys.argv) > 3 else str(ROOT / "deliverables/REFERENCES_corrected.docx")
NEWD = Path(sys.argv[4] if len(sys.argv) > 4 else ROOT / "source/newrefs")
AUDIT = ROOT / "deliverables/REFERENCES_audit.csv"
MAXAU = 30                                                     # as in REFERENCES_corrected

# ---------------------------------------------------------------- new references
def pages(p):
    """'980-7' -> '980–987'; e-pages and article numbers unchanged."""
    m = re.fullmatch(r"(\d+)-(\d+)", p or "")
    if not m:
        return p
    a, b = m.groups()
    if len(b) < len(a):
        b = a[:len(a) - len(b)] + b
    return f"{a}–{b}"


def authors(names):
    s = ", ".join(names[:MAXAU])
    return s + ", et al" if len(names) > MAXAU else s


# PubMed titles kept in title case by the source journal, set in sentence case like the rest of the list
SENTENCE_CASE = {"Averaging Strategy To Reduce Variability in Target-Decoy Estimates of False Discovery Rate":
                 "Averaging strategy to reduce variability in target-decoy estimates of false discovery rate"}


def entry(au, title, journal, year, vol, pg):
    title = title.strip()
    title = SENTENCE_CASE.get(title.rstrip("."), title)
    if not title.endswith((".", "?", "!")):
        title += "."
    return f"{authors(au)}. {title} {journal}. {year};{vol}:{pg}."


def pm_authors(art):
    au = []
    for a in art.findall(".//AuthorList/Author"):
        if a.findtext("CollectiveName"):
            au.append(a.findtext("CollectiveName"))
        else:
            au.append(f"{a.findtext('LastName')} {a.findtext('Initials') or ''}".strip())
    return au


# full author lists for old entries the GM list cut to "First A, et al." (efetch_etal.xml)
AUTH = {art.findtext(".//MedlineCitation/PMID"): pm_authors(art)
        for f in sorted(NEWD.glob("efetch*.xml")) for art in ET.parse(f).getroot().iter("PubmedArticle")}

NEW = {}                                                       # "Goeman 2007" -> dict
summ = json.load(open(NEWD / "esummary.json"))["result"]
xml = ET.parse(NEWD / "efetch.xml").getroot()
for art in xml.iter("PubmedArticle"):
    pmid = art.findtext(".//MedlineCitation/PMID")
    s = summ[pmid]
    au = pm_authors(art)
    year = s["pubdate"][:4]
    pg = s["pages"] or re.sub(r"^pii:\s*|\.?\s*doi:.*$", "", s.get("elocationid", "")).strip()
    jt = art.findtext(".//MedlineJournalInfo/MedlineTA")
    title = "".join(art.find(".//ArticleTitle").itertext())
    key = f"{au[0].rsplit(' ', 1)[0]} {year}"
    NEW[key] = dict(pmid=pmid, text=entry(au, title, jt, year, s["volume"], pages(pg)))

JABBR = {"The Annals of Applied Statistics": "Ann Appl Stat", "The Annals of Statistics": "Ann Stat"}
for f in sorted(NEWD.glob("crossref_*.json")):
    m = json.load(open(f))["message"]
    au = [f"{a['family']} {''.join(w[0] for w in re.split(r'[\s.-]+', a['given']) if w)}" for a in m["author"]]
    year = str(m["issued"]["date-parts"][0][0])
    key = f"{m['author'][0]['family']} {year}"
    NEW[key] = dict(pmid="", doi=m["DOI"], text=entry(au, m["title"][0], JABBR[m["container-title"][0]],
                                                     year, m["volume"], pages(m["page"])))

# ---------------------------------------------------------------- old references
cd = docx.Document(CORR)
OLD = {}
for p in cd.paragraphs:
    m = re.match(r"(\d+)\.\s+(.*)", p.text, re.S)
    if m:
        OLD[int(m.group(1))] = m.group(2).strip()
assert sorted(OLD) == list(range(1, len(OLD) + 1)), "corrected list is not 1..N"
OLDPMID = {r["pmid"]: int(r["ref"]) for r in csv.DictReader(open(AUDIT)) if r["pmid"].isdigit()}
etal = []
for pmid, n in OLDPMID.items():
    m = re.match(r"([^,.]+), et al\. ", OLD[n])
    if m and pmid in AUTH:
        assert AUTH[pmid][0].split()[0] == m.group(1).split()[0], (n, AUTH[pmid][0], m.group(1))
        OLD[n] = authors(AUTH[pmid]) + ". " + OLD[n][m.end():]
        etal.append(n)

# ---------------------------------------------------------------- the document
d = docx.Document(SRC)
body = d.element.body


H2 = {s.style_id for s in d.styles if s.name == "Heading 2"}     # the id is localised ("4" here)


def is_h2(p):
    st = p.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return st is not None and st.get(qn("w:val")) in H2


def ptext(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


# reading order: every paragraph outside the reference list, table cells cell by cell
order, reflist, in_refs = [], [], False
for el in body.iterchildren():
    if el.tag == qn("w:p"):
        if is_h2(el):
            in_refs = ptext(el).strip() == "References"
            if in_refs:
                refhead = el
            continue
        (reflist if in_refs else order).append(el)
    elif el.tag == qn("w:tbl"):
        order.extend(el.iter(qn("w:p")))
assert reflist, "no References section found"

NUMTOK = r"\[\s*\d{1,3}(?:\s*[–-]\s*\d{1,3})?(?:\s*,\s*\d{1,3}(?:\s*[–-]\s*\d{1,3})?)*\s*\]"
PHTOK = r"\{[^{}\[\]]+?\}"
TOK = f"(?:{NUMTOK}|{PHTOK})"
CIT = re.compile(f"{TOK}(?:\\s*{TOK})*")
PHPART = re.compile(r"^\s*(\S.*?)\s+(\d{4})\s*$")
unparsed = []


def parse(s):
    """Keys of one merged citation: ('old', n) or ('new', 'Author Year'); None if not a citation."""
    keys = []
    for tok in re.findall(TOK, s):
        if tok.startswith("["):
            for part in tok[1:-1].split(","):
                a, _, b = re.sub(r"\s", "", part).replace("–", "-").partition("-")
                lo, hi = int(a), int(b or a)
                if not (1 <= lo <= hi <= max(OLD)):
                    return None
                keys += [("old", n) for n in range(lo, hi + 1)]
        else:
            for part in tok[1:-1].split(";"):
                m = PHPART.match(part)
                if not m:
                    return None
                k = f"{m.group(1)} {m.group(2)}"
                k = unicodedata.normalize("NFC", k)
                if k not in NEW:
                    raise SystemExit(f"placeholder {{{k}}} has no fetched record in {NEWD}")
                old = OLDPMID.get(NEW[k]["pmid"])
                keys.append(("old", old) if old else ("new", k))
    return keys


hits = []                                                      # (paragraph, start, end, keys)
for p in order:
    t = ptext(p)
    for m in CIT.finditer(t):
        keys = parse(m.group(0))
        if keys is None:
            unparsed.append(m.group(0))
            continue
        hits.append((p, m.start(), m.end(), keys))

NUM = {}
for _, _, _, keys in hits:
    for k in keys:
        NUM.setdefault(k, len(NUM) + 1)


def fmt(nums):
    nums = sorted(set(nums))
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i >= 2:
            out.append(f"{nums[i]}–{nums[j]}")
        else:
            out += [str(n) for n in nums[i:j + 1]]
        i = j + 1
    return "[" + ", ".join(out) + "]"


def replace(p, start, end, new):
    """Replace characters [start, end) of the paragraph's text, editing w:t in place."""
    pos = 0
    for t in list(p.iter(qn("w:t"))):
        s = t.text or ""
        a, b = pos, pos + len(s)
        pos = b
        if b <= start or a >= end:
            continue
        lo, hi = max(start, a) - a, min(end, b) - a
        t.text = s[:lo] + (new if a <= start < b else "") + s[hi:]
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


for p, s, e, keys in reversed(hits):                           # right to left keeps offsets valid
    replace(p, s, e, fmt(NUM[k] for k in keys))

# ---------------------------------------------------------------- the list
tmpl_first, tmpl_rest = reflist[0], reflist[1] if len(reflist) > 1 else reflist[0]
for p in reflist:
    body.remove(p)


def make(tmpl, text):
    p = copy.deepcopy(tmpl)
    runs = [r for r in p.iter(qn("w:r"))]
    keep = runs[0]
    for r in runs[1:]:
        r.getparent().remove(r)
    for c in list(keep):
        if c.tag != qn("w:rPr"):
            keep.remove(c)
    t = keep.makeelement(qn("w:t"), {})
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    keep.append(t)
    return p


anchor = refhead
for k, n in sorted(NUM.items(), key=lambda kv: kv[1]):
    text = OLD[k[1]] if k[0] == "old" else NEW[k[1]]["text"]
    p = make(tmpl_first if n == 1 else tmpl_rest, f"{n}. {text}")
    anchor.addnext(p)
    anchor = p

d.save(OUT)

# ---------------------------------------------------------------- checks and report
d2 = docx.Document(OUT)
alltext = "\n".join(ptext(p) for p in d2.element.body.iter(qn("w:p")))
left = re.findall(r"\{[^{}]*\d{4}[^{}]*\}", alltext)
assert not left, f"placeholders left: {left}"
cited = set()
for p in d2.element.body.iter(qn("w:p")):
    if re.match(r"\d+\.\s", ptext(p)):
        continue
    for m in re.finditer(NUMTOK, ptext(p)):
        for part in m.group(0)[1:-1].split(","):
            a, _, b = re.sub(r"\s", "", part).replace("–", "-").partition("-")
            if a.isdigit() and int(a) <= len(NUM):
                cited |= set(range(int(a), int(b or a) + 1))
assert cited == set(range(1, len(NUM) + 1)), f"uncited: {set(range(1, len(NUM) + 1)) - cited}"

dropped = sorted(set(OLD) - {k[1] for k in NUM if k[0] == "old"})
print(f"references {len(NUM)}: old kept {sum(k[0] == 'old' for k in NUM)}, new {sum(k[0] == 'new' for k in NUM)}; "
      f"dropped old {dropped}")
print("old entries given full author lists (were 'First, et al.'):", sorted(etal))
print("citations rewritten:", len(hits), "| not parsed as citations:", unparsed or "none")
print("old -> new:", ", ".join(f"{k[1]}->{n}" for k, n in NUM.items() if k[0] == "old"))
print("new:", "; ".join(f"{k[1]} -> {n}" for k, n in NUM.items() if k[0] == "new"))
for k, n in NUM.items():
    if k[0] == "new":
        print(f"  {n}. {NEW[k[1]]['text']}")
