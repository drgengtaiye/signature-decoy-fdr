"""Code availability for the Genome Biology manuscript.

usage: build_manuscript_GB_v3.py <MANUSCRIPT_GB_v2.docx> <out.docx>

The analysis code is released as a public repository holding exactly Additional file 3; the
earlier immunoTriage-paper repository (which also holds the separate paper's analyses) and the
immunoTriage R package (which this paper's code does not call) are no longer cited. Reference 99
was cited only here and is the last reference, so no renumbering is needed.

The figure-legend section also carried the GM figures as embedded images (old Figs 1-8, including the
old Fig. 2 with its triage rings and old Figs 7-8 of the separate paper, and no Fig. 4). Figures are
uploaded as separate files, so every embedded image goes and the legends stay.

Last, copyedits found on proofreading: the first author's initials in the contributions, title-page
and funding spacing, CRediT role names, American spelling outside the reference list, and three
abbreviations missing from the list.
"""
import re
import sys

import docx
from docx.oxml.ns import qn

SRC, OUT = sys.argv[1:3]
REPO = "https://github.com/drgengtaiye/signature-decoy-fdr"
ACCESSED = "20 Oct 2026"
d = docx.Document(SRC)
ps = d.paragraphs
find = lambda start: [p for p in ps if p.text.startswith(start)]


def set_text(p, new):
    """p has one run; replace its text and keep the run's format."""
    assert len(p.runs) == 1, p.text[:60]
    p.runs[0].text = new


(avail,) = find("All analysis code is provided as Additional file 3")
set_text(avail, "All analysis code is provided as Additional file 3 and in the signature-decoy-fdr repository [98], "
         "including the scripts that score each signature by its publication’s own method against its size- and "
         "engine-matched null (null_matched_v2.R, null_matched_prog.R) and those that grow and score the decoy "
         "modules and estimate the false discovery rate (calib_decoys_v3.R, calib_score_v3.R, null_natpool_v4.R, "
         "calib_v3_fdr.py, calib_v3_robust.py). The repository will be archived with a DOI on acceptance.")
(r98,) = find("98. immunoTriage-paper")
set_text(r98, f"98. signature-decoy-fdr: analysis code for this study. {REPO}. Accessed {ACCESSED}.")
(r99,) = find("99. immunoTriage")
r99._p.getparent().remove(r99._p)


# ---- copyedits found on proofreading (2026-10-20)
def run_sub(p, old, new):
    """Replace old by new inside one run of paragraph p."""
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return
    raise AssertionError((old, p.text[:60]))


ps = d.paragraphs
# title page: "Zheng3,  Xuetao" and "Peng3,Jun"
authors = ps[1]
run_sub(authors, " Xuetao Lei", "Xuetao Lei")
k = [i for i, r in enumerate(authors.runs) if r.text == "Jun Qu"][0]
assert authors.runs[k - 1].text == ","
authors.runs[k - 1].text = ", "
# authors' contributions: the first author's initials lost a letter; CRediT role names; spacing
(contrib,) = [ps[i + 1] for i, p in enumerate(ps) if p.text.strip() in ("Authors’ contributions", "Authors' contributions")]
assert contrib.runs[0].text == "L."
contrib.runs[0].text = "W.L."
run_sub(contrib, " Z.Y.", "Z.Y.")
for r in contrib.runs:
    r.text = (r.text.replace("Writing- original draft", "Writing – original draft")
              .replace("Writing-original draft", "Writing – original draft")
              .replace("Writing-review & editing", "Writing – review & editing"))
(fund,) = [ps[i + 1] for i, p in enumerate(ps) if p.text.strip() == "Funding"]
run_sub(fund, "Hospital  (", "Hospital (")
(fig4ac,) = [p for p in ps if "(Fig. 4a,c)" in p.text]
run_sub(fig4ac, "Fig. 4a,c", "Fig. 4a, c")

# American spelling in the text (reference titles keep their published spelling)
US = {"favourable": "favorable", "favoured": "favored", "centred": "centered", "centre": "center",
      "standardisation": "standardization", "standardised": "standardized", "modelled": "modeled",
      "tumour": "tumor"}
rx_us = re.compile(r"\b(" + "|".join(US) + r")\b")
texts = [p.text for p in ps]
r0, r1 = texts.index("References"), texts.index("Figure titles and legends")
cells = [cp for tb in d.tables for row in tb.rows for c in row.cells for cp in c.paragraphs]
for p in [p for i, p in enumerate(ps) if not r0 < i < r1] + cells:
    for r in p.runs:
        if rx_us.search(r.text):
            r.text = rx_us.sub(lambda m: US[m.group(1)], r.text)

# abbreviations used in formulas and figures but not listed
(abbr,) = [ps[i + 1] for i, p in enumerate(ps) if p.text.strip() == "Abbreviations"]
lines = [l for l in abbr.text.split("\n") if l.strip()]
lines += ["C-index: concordance index", "pCR: pathological complete response", "PFS: progression-free survival"]
assert len(abbr.runs) == 1
abbr.runs[0].text = "\n".join(sorted(set(lines), key=str.lower))

# embedded figures: drop each drawing, and its paragraph when nothing else is in it
for dr in list(d.element.body.iter(qn("w:drawing"))):
    p = next(a for a in dr.iterancestors() if a.tag == qn("w:p"))
    run = dr.getparent()
    run.remove(dr)
    if not "".join(t.text or "" for t in p.iter(qn("w:t"))).strip() and not list(p.iter(qn("w:drawing"))):
        p.getparent().remove(p)
for rid, rel in list(d.part.rels.items()):
    if rel.reltype.endswith("/image"):
        d.part.drop_rel(rid)

d.save(OUT)
d2 = docx.Document(OUT)
assert not list(d2.element.body.iter(qn("w:drawing"))), "embedded image left"
assert not [r for r in d2.part.rels.values() if r.reltype.endswith("/image")], "image part left"
txt = "\n".join(p.text for p in d2.paragraphs)
assert "[99]" not in txt and "dryichen" not in txt and "immunoTriage" not in txt, "residual"
print("saved", OUT)
